"""Convert the two clan Excel workbooks into data/data.js for the site.

Usage (from the stray-alliance folder):
    python tools/convert.py

Looks for the workbooks one level above the repo folder by default;
pass paths explicitly to override:
    python tools/convert.py path/to/aPAW.xlsx path/to/SOC.xlsx
"""

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
DEFAULT_APAW = REPO.parent / "SurvivorIO-2 aPAWcalypse.xlsx"
DEFAULT_SOC = REPO.parent / "SurvivorIO-2 StraytOutofCompton.xlsx"

DATE_SHEET = re.compile(r"^\d{4}-\d{2}-\d{2}$")
# Old-mechanic headers like "1 (26,800)" recorded the boss level the
# opponent sent that day; the game changed mechanics so tracking stopped.
BOSS_LEVEL = re.compile(r"\(([\d,]+)\)")


def cell_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip().replace(",", "")
    return int(float(s)) if s and s.replace(".", "").isdigit() else None


HEADER_ALIASES = {
    "place": "place", "placement": "place",
    "tickets used": "used", "used": "used",
    "tickets remaining": "left", "tickets left": "left",
    "remaining": "left", "left": "left", "tickets": "left",
}


def read_results(ws):
    """Parse the Results sheet: per-week placement plus optional ticket columns.

    Columns: A=date, B=place. Without a header row, the ticket columns are
    found automatically — the first two columns after Place that contain
    numbers (columns holding dates, like SOC's recorded-on column, are
    skipped) are Tickets Used and Tickets Remaining, left to right.
    A header row naming the columns (see HEADER_ALIASES, any order)
    overrides the auto-detection.
    """
    rows = list(ws.iter_rows(values_only=True))
    cols = {"place": 1, "used": None, "left": None}
    header_found = False
    for row in rows:
        header = {str(c).strip().lower(): i for i, c in enumerate(row) if isinstance(c, str)}
        hit = {HEADER_ALIASES[k]: i for k, i in header.items() if k in HEADER_ALIASES}
        if hit:
            cols.update(hit)
            header_found = True
            break

    if not header_found:
        numeric_cols = []
        width = max((len(r) for r in rows), default=0)
        for i in range(2, width):
            vals = [r[i] for r in rows if len(r) > i and r[i] is not None]
            if not vals or any(isinstance(v, (datetime, date)) for v in vals):
                continue
            if any(cell_num(v) is not None for v in vals):
                numeric_cols.append(i)
        cols["used"] = numeric_cols[0] if numeric_cols else None
        cols["left"] = numeric_cols[1] if len(numeric_cols) > 1 else None

    def cell(row, key):
        i = cols.get(key)
        return row[i] if i is not None and len(row) > i else None

    results = {}
    for row in rows:
        d = row[0]
        if not isinstance(d, (datetime, date)):
            continue
        place = cell(row, "place")
        results[d.strftime("%Y-%m-%d")] = {
            "placement": str(place).strip() if place else None,
            "ticketsUsed": cell_num(cell(row, "used")),
            "ticketsLeft": cell_num(cell(row, "left")),
        }
    return results


def read_clan(path, clan_id):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    results = {}
    if "Results" in wb.sheetnames:
        results = read_results(wb["Results"])

    weeks = []
    for name in wb.sheetnames:
        if not DATE_SHEET.match(name):
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[1] if len(rows) > 1 else ()

        boss_levels = []
        for c in range(2, 5):
            m = BOSS_LEVEL.search(str(header[c])) if len(header) > c and header[c] else None
            boss_levels.append(int(m.group(1).replace(",", "")) if m else None)

        # Row 1 cells above the day headers hold free-text result notes,
        # e.g. "2-0" or "Cleared, overloaded opponent". Excel silently turns
        # "1-1" into the date Jan 1, so dates are mapped back to month-day.
        row1 = rows[0] if rows else ()
        day_notes = []
        for c in range(2, 5):
            v = row1[c] if len(row1) > c else None
            if isinstance(v, (datetime, date)):
                v = f"{v.month}-{v.day}"
            s = str(v).strip() if v is not None else ""
            day_notes.append(s if s else None)

        players = []
        for row in rows[2:]:
            if not row or row[0] is None or str(row[0]).strip() == "":
                continue
            players.append({
                "name": str(row[0]).strip(),
                "p1": cell_num(row[1]) if len(row) > 1 else None,
                "days": [cell_num(row[c]) if len(row) > c else None for c in (2, 3, 4)],
            })

        n = len(players)
        # A battle day left blank for (almost) the whole roster was a blowout
        # week where hits weren't required, so scores were never recorded.
        day_tracked = []
        for d in range(3):
            blanks = sum(1 for p in players if p["days"][d] is None)
            day_tracked.append(not (n > 0 and blanks >= n - 1))

        res = results.get(name, {})
        weeks.append({
            "date": name,
            "placement": res.get("placement"),
            "ticketsUsed": res.get("ticketsUsed"),
            "ticketsLeft": res.get("ticketsLeft"),
            "dayNotes": day_notes,
            "dayTracked": day_tracked,
            "bossLevels": boss_levels,  # legacy: boss level faced, early weeks only
            "players": players,
        })

    wb.close()
    weeks.sort(key=lambda w: w["date"])
    # current ticket count = most recent week with a remaining count recorded
    tickets = next((w["ticketsLeft"] for w in reversed(weeks) if w["ticketsLeft"] is not None), None)
    return {"id": clan_id, "weeks": weeks, "tickets": tickets}


def main():
    apaw_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_APAW
    soc_path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_SOC

    data = {
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "clans": {
            "apaw": read_clan(apaw_path, "apaw"),
            "soc": read_clan(soc_path, "soc"),
        },
    }

    out = REPO / "data" / "data.js"
    out.parent.mkdir(exist_ok=True)
    out.write_text(
        "// Generated by tools/convert.py — do not edit by hand.\n"
        "window.CLAN_DATA = " + json.dumps(data, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    for cid, clan in data["clans"].items():
        print(f"{cid}: {len(clan['weeks'])} weeks, "
              f"latest {clan['weeks'][-1]['date']} "
              f"({len(clan['weeks'][-1]['players'])} players)")
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
